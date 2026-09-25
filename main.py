from pathlib import Path
import json
import re
import uuid
from collections import Counter
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from pydantic import BaseModel

from graph_analysis import (
    build_graph,
    calculate_metrics,
    detect_communities,
    detect_suspicious_patterns
)

from database import (
    init_database,
    create_case,
    get_cases,
    get_case,
    update_case_status,
    delete_case,
    save_case_entities,
    get_case_entities,
    get_case_relationships,
    save_previous_case,
    get_previous_cases,
    get_previous_case,
    save_previous_case_link
)

from fir_analysis import analyze_fir_text

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    import cv2
except ImportError:
    cv2 = None

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
UPLOAD_DIR = BASE_DIR / "uploads"
VIDEO_DIR = UPLOAD_DIR / "cctv"
UPLOAD_DIR.mkdir(exist_ok=True)
VIDEO_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI(
    title="AI Criminal Network Analysis System",
    description="AI-powered investigator-assistance prototype using synthetic demonstration data",
    version="2.2.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"]
)

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
app.mount("/uploads", StaticFiles(directory=UPLOAD_DIR), name="uploads")
init_database()

class CaseCreate(BaseModel):
    case_id: str
    title: str
    description: str = ""
    priority: str = "MEDIUM"

class CaseStatusUpdate(BaseModel):
    status: str

class FIRAnalyzeRequest(BaseModel):
    text: str
    create_case: bool = False
    priority: str = "MEDIUM"


def _read_uploaded_document(data: bytes, filename: str, content_type: str) -> tuple[str, str]:
    """Read PDF/DOCX/TXT uploads using the same parser as FIR intelligence."""
    filename_l = (filename or "").lower()
    content_type_l = (content_type or "").lower()
    if not data:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File is too large. Maximum size is 20 MB")
    try:
        from io import BytesIO
        if filename_l.endswith(".pdf") or content_type_l == "application/pdf":
            if PdfReader is None:
                raise HTTPException(status_code=500, detail="PDF support is not installed.")
            reader = PdfReader(BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
            doc_type = "PDF"
        elif filename_l.endswith(".docx") or content_type_l == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            if Document is None:
                raise HTTPException(status_code=500, detail="DOCX support is not installed.")
            doc = Document(BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip()).strip()
            doc_type = "DOCX"
        elif filename_l.endswith(".txt") or content_type_l.startswith("text/"):
            text = data.decode("utf-8-sig", errors="replace").strip()
            doc_type = "TXT"
        else:
            raise HTTPException(status_code=415, detail="Unsupported file type. Upload PDF, DOCX, or TXT.")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read the document: {exc}")
    if not text:
        raise HTTPException(status_code=400, detail="No readable text was found in the document.")
    return text, doc_type

def _norm_value(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()

_PATTERN_GROUPS = {
    "theft": {"theft","stolen","stealing","burglary","robbery","robbed","snatching","larceny"},
    "assault": {"assault","attack","attacked","injury","fight","beating"},
    "fraud": {"fraud","scam","cheating","forgery","forged","financial"},
    "cyber": {"cyber","online","phishing","malware","hacking","account takeover","otp"},
    "drugs": {"drug","narcotic","contraband","smuggling","cocaine","heroin"},
    "weapons": {"weapon","firearm","gun","pistol","ammunition"},
    "kidnapping": {"kidnap","abduction","abducted","hostage"},
    "traffic": {"vehicle","collision","accident","hit and run","traffic"}
}

def _pattern_tags(text):
    s = _norm_value(text)
    return sorted([k for k, words in _PATTERN_GROUPS.items() if any(w in s for w in words)])

def _case_similarity(current_result, historical_result, current_case=None, historical_text=""):
    c = current_result.get("extracted", {}) or {}
    h = historical_result.get("extracted", {}) or {}
    reasons, score = [], 0.0

    def normset(key):
        return {_norm_value(x) for x in (c.get(key, []) or []) if _norm_value(x)}

    def histset(key):
        return {_norm_value(x) for x in (h.get(key, []) or []) if _norm_value(x)}

    common_people = sorted(normset("persons") & histset("persons"))
    common_locations = sorted(normset("locations") & histset("locations"))
    common_vehicles = sorted(normset("vehicles") & histset("vehicles"))
    common_phones = sorted(normset("phones") & histset("phones"))
    common_accounts = sorted(normset("accounts") & histset("accounts"))

    if common_people:
        score += 45
        reasons.append({"type":"PERSON","label":"Same mentioned person/entity","matches":common_people[:10]})
    if common_locations:
        score += 30
        reasons.append({"type":"LOCATION","label":"Same mentioned location","matches":common_locations[:10]})
    if common_vehicles:
        score += 10
        reasons.append({"type":"VEHICLE","label":"Same vehicle reference","matches":common_vehicles[:10]})
    if common_phones:
        score += 8
        reasons.append({"type":"PHONE","label":"Same communication identifier","matches":common_phones[:10]})
    if common_accounts:
        score += 8
        reasons.append({"type":"ACCOUNT","label":"Same financial identifier","matches":common_accounts[:10]})

    ct = set(_pattern_tags(
        " ".join([current_result.get("title",""), current_case.get("description","") if current_case else "",
                  current_result.get("case_id",""), " ".join(c.get("persons",[])), " ".join(c.get("locations",[])),
                  current_result.get("note","")])
    ))
    ht = set(_pattern_tags(" ".join([historical_result.get("title",""), historical_text,
                                      " ".join(h.get("persons",[])), " ".join(h.get("locations",[]))])))
    common_patterns = sorted(ct & ht)
    if common_patterns:
        score += min(25, 12 + 4*len(common_patterns))
        reasons.append({"type":"PATTERN","label":"Similar incident pattern indicators","matches":common_patterns})

    # Cap at 100 and provide a reasoned score rather than a claim of certainty.
    score = min(100.0, round(score, 1))
    return {"score": score, "reasons": reasons}


@app.get("/")
def home():
    return FileResponse(STATIC_DIR / "index.html")

@app.get("/health")
def health():
    return {"status": "healthy", "system": "AI Criminal Network Analysis System"}

@app.get("/api/network")
def network():
    graph = build_graph()
    communities = detect_communities(graph)
    return {"entities": graph.number_of_nodes(), "relationships": graph.number_of_edges(), "communities": len(set(communities.values())), "status": "ANALYZED"}

@app.get("/api/top-entities")
def top_entities():
    graph = build_graph()
    metrics = calculate_metrics(graph)
    return {"entities": metrics[:10]}

@app.get("/api/entities")
def entities():
    graph = build_graph()
    result = []
    for node, data in graph.nodes(data=True):
        result.append({"id": node, "type": data.get("node_type", "unknown"), "name": data.get("name", node), "city": data.get("city", "")})
    return {"entities": result}

@app.get("/api/communities")
def communities():
    graph = build_graph()
    community_map = detect_communities(graph)
    result = {}
    for node, community in community_map.items():
        result.setdefault(community, []).append(node)
    return {"communities": result}

@app.get("/api/graph")
def graph_data():
    graph = build_graph()
    nodes = [{"id": node, "label": data.get("name", node), "type": data.get("node_type", "unknown")} for node, data in graph.nodes(data=True)]
    edges = [{"source": source, "target": target, "relationship": data.get("relationship", "UNKNOWN")} for source, target, data in graph.edges(data=True)]
    return {"nodes": nodes, "edges": edges}


@app.get("/api/map-data")
def map_data(case_id: str | None = None):
    """Return case-specific map points using public city-center coordinates only."""
    # Be tolerant of older frontend dropdowns that accidentally sent
    # "<case_id> Case" instead of the actual case_id.
    if case_id:
        case_id = case_id.strip()
        if case_id.endswith(" Case") and not get_case(case_id):
            candidate = case_id[:-5].strip()
            if get_case(candidate):
                case_id = candidate
    city_coords = {
        "eluru": (16.7107, 81.0952),
        "vijayawada": (16.5062, 80.6480),
        "visakhapatnam": (17.6868, 83.2185),
        "vizag": (17.6868, 83.2185),
        "hyderabad": (17.3850, 78.4867),
        "secunderabad": (17.4399, 78.4983),
        "guntur": (16.3067, 80.4365),
        "rajahmundry": (17.0005, 81.8040),
        "tirupati": (13.6288, 79.4192),
        "warangal": (17.9784, 79.5941),
    }
    def resolve_city(name):
        text = str(name or "").strip().lower()
        matches = [(c, xy) for c, xy in city_coords.items() if c in text]
        return max(matches, key=lambda x: len(x[0])) if matches else None

    if case_id:
        if not get_case(case_id):
            raise HTTPException(status_code=404, detail="Case not found")
        entities = get_case_entities(case_id)
        relationships = get_case_relationships(case_id)
        by_id = {str(e["id"]): e for e in entities}
        locations = []
        for e in entities:
            if str(e.get("type", "")).upper() != "LOCATION":
                continue
            resolved = resolve_city(e.get("name", ""))
            if not resolved:
                continue
            city, (lat, lon) = resolved
            connections = []
            for rel in relationships:
                src, dst = str(rel.get("source")), str(rel.get("target"))
                if src == str(e["id"]): other_id = dst
                elif dst == str(e["id"]): other_id = src
                else: continue
                other = by_id.get(other_id, {})
                connections.append({"id": other_id, "name": other.get("name", other_id), "type": other.get("type", "UNKNOWN"), "relationship": rel.get("relationship", "RELATED")})
            locations.append({"id": e["id"], "name": e.get("name", e["id"]), "city": city.title(), "lat": lat, "lon": lon, "connections": connections, "coordinate_type": "PUBLIC_CITY_CENTER"})
        return {"case_id": case_id, "locations": locations, "note": "Real public city-center coordinates are used only for visualization; they do not represent precise incident addresses."}

    graph = build_graph()
    locations = []
    for node, data in graph.nodes(data=True):
        if str(data.get("node_type", "")).lower() != "location": continue
        resolved = resolve_city(data.get("name", ""))
        if not resolved: continue
        city, (lat, lon) = resolved
        connections = [{"id": n, "name": graph.nodes[n].get("name", n), "type": graph.nodes[n].get("node_type", "unknown"), "relationship": graph.edges[node, n].get("relationship", "RELATED")} for n in graph.neighbors(node)]
        locations.append({"id": node, "name": data.get("name", node), "city": city.title(), "lat": lat, "lon": lon, "connections": connections, "coordinate_type": "PUBLIC_CITY_CENTER"})
    return {"case_id": None, "locations": locations, "note": "Real public city-center coordinates are used only for visualization."}

@app.get("/api/cases/{case_id}/entities/{entity_id}")
def case_entity_details(case_id: str, entity_id: str):
    if not get_case(case_id):
        raise HTTPException(status_code=404, detail="Case not found")
    entities = get_case_entities(case_id)
    entity = next((e for e in entities if str(e["id"]) == str(entity_id)), None)
    if not entity:
        raise HTTPException(status_code=404, detail="Entity not found in this case")
    relationships = get_case_relationships(case_id)
    by_id = {str(e["id"]): e for e in entities}
    connections = []
    for r in relationships:
        src, dst = str(r["source"]), str(r["target"])
        if src == str(entity_id):
            other_id = dst
        elif dst == str(entity_id):
            other_id = src
        else:
            continue
        other = by_id.get(other_id, {})
        connections.append({
            "id": other_id,
            "name": other.get("name", other_id),
            "type": other.get("type", "UNKNOWN"),
            "relationship": r.get("relationship", "RELATED")
        })
    return {"case_id": case_id, "id": entity_id, "details": entity, "connections": connections}

@app.get("/api/entity/{entity_id}")
def entity_details(entity_id: str):
    graph = build_graph()
    if entity_id not in graph:
        raise HTTPException(status_code=404, detail="Entity not found")
    data = graph.nodes[entity_id]
    connections = [{"id": n, "type": graph.nodes[n].get("node_type", "unknown")} for n in graph.neighbors(entity_id)]
    return {"id": entity_id, "details": dict(data), "connections": connections}

@app.get("/api/alerts")
def alerts():
    graph = build_graph()
    return {"alerts": detect_suspicious_patterns(graph)}

@app.post("/api/cases")
def create_new_case(case: CaseCreate):
    existing = get_case(case.case_id)
    if existing:
        raise HTTPException(status_code=409, detail="Case ID already exists")
    if not case.case_id.strip():
        raise HTTPException(status_code=400, detail="Case ID cannot be empty")
    if not case.title.strip():
        raise HTTPException(status_code=400, detail="Case title cannot be empty")
    if case.priority.upper() not in {"LOW", "MEDIUM", "HIGH", "CRITICAL"}:
        raise HTTPException(status_code=400, detail="Invalid priority")
    new_case = create_case(case.case_id.strip(), case.title.strip(), case.description.strip(), case.priority.upper())
    return {"message": "Case created successfully", "case": new_case}

@app.get("/api/cases")
def list_cases():
    return {"cases": get_cases()}

@app.get("/api/cases/{case_id}")
def single_case(case_id: str):
    case = get_case(case_id)
    if not case:
        raise HTTPException(status_code=404, detail="Case not found")
    return {"case": case}

@app.delete("/api/cases/{case_id}")
def remove_case(case_id: str):
    deleted = delete_case(case_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="Case not found")
    return {"message": "Case deleted successfully", "case_id": case_id}

@app.patch("/api/cases/{case_id}/status")
def change_case_status(case_id: str, update: CaseStatusUpdate):
    status = update.status.upper()
    if status not in {"OPEN", "UNDER REVIEW", "CLOSED", "ON HOLD"}:
        raise HTTPException(status_code=400, detail="Invalid case status")
    updated = update_case_status(case_id, status)
    if not updated:
        raise HTTPException(status_code=404, detail="Case not found")
    return {"message": "Case status updated", "case": updated}

def _slug_entity_id(case_id, entity_type, name, index):
    slug=re.sub(r"[^A-Za-z0-9]+", "-", str(name)).strip("-").upper()[:28] or "ENTITY"
    return f"FIR-{case_id}-{entity_type}-{index}-{slug}"

def _build_case_entities(case_id, result):
    extracted=result.get("extracted", {}) or {}
    matched={str(m.get("name", "")).strip().lower():m for m in (result.get("matched_entities", []) or []) if m.get("name")}
    entities=[]; id_by_name={}; counts={}
    type_map={"persons":"PERSON","locations":"LOCATION","vehicles":"VEHICLE","phones":"PHONE","accounts":"ACCOUNT"}
    for key, etype in type_map.items():
        for value in extracted.get(key, []) or []:
            name=str(value).strip()
            if not name: continue
            norm=name.lower()
            if norm in id_by_name: continue
            counts[etype]=counts.get(etype,0)+1
            m=matched.get(norm)
            entity_id=str(m.get("id")) if m and m.get("id") else _slug_entity_id(case_id,etype,name,counts[etype])
            entity={"id":entity_id,"name":str(m.get("name",name)) if m else name,"type":str(m.get("type",etype)).upper() if m else etype,"city":"","source":"FIR","metadata":json.dumps({"fir_number":result.get("fir_number",""),"date":result.get("date","")})}
            entities.append(entity); id_by_name[norm]=entity_id
    rels=[]
    for r in result.get("relationships",[]) or []:
        src=id_by_name.get(str(r.get("source","")).strip().lower()); dst=id_by_name.get(str(r.get("target","")).strip().lower())
        if src and dst and src!=dst: rels.append({"source":src,"target":dst,"relationship":r.get("relationship","RELATED")})
    return entities, rels

@app.post("/api/cases/{case_id}/entities")
def attach_case_entities(case_id: str, payload: dict):
    if not get_case(case_id): raise HTTPException(status_code=404, detail="Case not found")
    entities, relationships=_build_case_entities(case_id,payload)
    save_case_entities(case_id,entities,relationships)
    return {"message":"Case entities saved","case_id":case_id,"entities":entities,"relationships":relationships,"entity_count":len(entities),"relationship_count":len(relationships)}

@app.get("/api/cases/{case_id}/entities")
def list_case_entities(case_id: str):
    if not get_case(case_id): raise HTTPException(status_code=404, detail="Case not found")
    return {"case_id":case_id,"entities":get_case_entities(case_id)}

@app.get("/api/cases/{case_id}/graph")
def case_graph(case_id: str):
    if not get_case(case_id): raise HTTPException(status_code=404, detail="Case not found")
    entities=get_case_entities(case_id); relationships=get_case_relationships(case_id)
    return {"case_id":case_id,"nodes":[{"id":e["id"],"label":e["name"],"type":e["type"]} for e in entities],"edges":relationships}

@app.get("/api/cases/{case_id}/analysis")
def case_analysis(case_id: str):
    """Return explainable, case-scoped network analysis for synthetic investigation data."""
    if not get_case(case_id):
        raise HTTPException(status_code=404, detail="Case not found")
    entities = get_case_entities(case_id)
    relationships = get_case_relationships(case_id)
    g = __import__("networkx").Graph()
    for e in entities:
        g.add_node(str(e["id"]), name=e.get("name", e["id"]), type=e.get("type", "UNKNOWN"))
    for r in relationships:
        src, dst = str(r.get("source")), str(r.get("target"))
        if src in g and dst in g and src != dst:
            g.add_edge(src, dst, relationship=r.get("relationship", "RELATED"))

    type_counts = {}
    for e in entities:
        t = str(e.get("type", "UNKNOWN")).upper()
        type_counts[t] = type_counts.get(t, 0) + 1

    relationship_counts = {}
    for r in relationships:
        rel = str(r.get("relationship", "RELATED")).upper()
        relationship_counts[rel] = relationship_counts.get(rel, 0) + 1

    if g.number_of_nodes():
        degree = dict(g.degree())
        betweenness = __import__("networkx").betweenness_centrality(g) if g.number_of_nodes() > 1 else {n: 0 for n in g.nodes()}
        ranked = sorted(g.nodes(), key=lambda n: (degree.get(n, 0), betweenness.get(n, 0)), reverse=True)
        top_entities = [
            {
                "id": n,
                "name": g.nodes[n].get("name", n),
                "type": g.nodes[n].get("type", "UNKNOWN"),
                "connections": degree.get(n, 0),
                "betweenness": round(betweenness.get(n, 0), 3)
            }
            for n in ranked[:8]
        ]
        components = list(__import__("networkx").connected_components(g))
        component_sizes = sorted((len(c) for c in components), reverse=True)
        density = round(__import__("networkx").density(g), 3) if g.number_of_nodes() > 1 else 0
    else:
        top_entities = []
        component_sizes = []
        density = 0

    return {
        "case_id": case_id,
        "entity_count": len(entities),
        "relationship_count": len(relationships),
        "density": density,
        "connected_components": len(component_sizes),
        "component_sizes": component_sizes,
        "entity_types": type_counts,
        "relationship_types": relationship_counts,
        "top_entities": top_entities,
        "status": "READY" if entities else "NO_CASE_ENTITIES",
        "note": "Synthetic case-scoped analysis for investigator review. Metrics are analytical aids and do not establish criminal activity or guilt."
    }


@app.get("/api/previous-cases")
def list_previous_case_files():
    return {"previous_cases": get_previous_cases()}

@app.get("/api/previous-cases/{previous_id}")
def previous_case_details(previous_id: int):
    row = get_previous_case(previous_id)
    if not row:
        raise HTTPException(status_code=404, detail="Previous case not found")
    try:
        row["analysis"] = json.loads(row.pop("extracted_json") or "{}")
    except Exception:
        row["analysis"] = {}
    row.pop("raw_text", None)
    return {"previous_case": row}

@app.post("/api/previous-cases/upload")
async def upload_previous_case(
    file: UploadFile = File(...),
    current_case_id: str = Form("")
):
    data = await file.read()
    text, doc_type = _read_uploaded_document(data, file.filename or "", file.content_type or "")
    result = analyze_fir_text(text, graph=build_graph())
    reference_id = "PC-" + uuid.uuid4().hex[:8].upper()
    row = save_previous_case(reference_id, file.filename or "previous-case", result, text)

    matches = []
    if current_case_id.strip():
        current_case = get_case(current_case_id.strip())
        if not current_case:
            raise HTTPException(status_code=404, detail="Current case not found")
        current_entities = get_case_entities(current_case_id.strip())
        # Use stored case entities when available; this preserves the current investigation context.
        current_result = {
            "case_id": current_case_id.strip(),
            "title": current_case.get("title",""),
            "extracted": {
                "persons": [e["name"] for e in current_entities if str(e.get("type","")).upper()=="PERSON"],
                "locations": [e["name"] for e in current_entities if str(e.get("type","")).upper()=="LOCATION"],
                "vehicles": [e["name"] for e in current_entities if str(e.get("type","")).upper()=="VEHICLE"],
                "phones": [e["name"] for e in current_entities if str(e.get("type","")).upper()=="PHONE"],
                "accounts": [e["name"] for e in current_entities if str(e.get("type","")).upper()=="ACCOUNT"],
            }
        }
        match = _case_similarity(current_result, result, current_case, text)
        if match["reasons"]:
            save_previous_case_link(row["id"], current_case_id.strip(), match["score"], match["reasons"])
            matches.append({
                "previous_case_id": row["id"],
                "reference_id": reference_id,
                "filename": file.filename,
                "title": result.get("title") or file.filename,
                "score": match["score"],
                "reasons": match["reasons"]
            })

    return {
        "message":"Previous case uploaded and analyzed",
        "reference_id":reference_id,
        "document_type":doc_type,
        "analysis":result,
        "matches":matches,
        "note":"Historical correlation is an investigative lead based on extracted synthetic data; it does not establish identity, guilt, or criminal involvement."
    }

@app.post("/api/previous-cases/match/{case_id}")
def match_previous_cases(case_id: str):
    current_case = get_case(case_id)
    if not current_case:
        raise HTTPException(status_code=404, detail="Case not found")
    current_entities = get_case_entities(case_id)
    current_result = {
        "case_id": case_id,
        "title": current_case.get("title",""),
        "extracted": {
            "persons": [e["name"] for e in current_entities if str(e.get("type","")).upper()=="PERSON"],
            "locations": [e["name"] for e in current_entities if str(e.get("type","")).upper()=="LOCATION"],
            "vehicles": [e["name"] for e in current_entities if str(e.get("type","")).upper()=="VEHICLE"],
            "phones": [e["name"] for e in current_entities if str(e.get("type","")).upper()=="PHONE"],
            "accounts": [e["name"] for e in current_entities if str(e.get("type","")).upper()=="ACCOUNT"],
        }
    }
    results=[]
    for row in get_previous_cases():
        full=get_previous_case(row["id"])
        try: hist=json.loads(full.get("extracted_json") or "{}")
        except Exception: hist={}
        match=_case_similarity(current_result,hist,current_case,full.get("raw_text",""))
        if match["reasons"]:
            save_previous_case_link(row["id"],case_id,match["score"],match["reasons"])
            results.append({"previous_case_id":row["id"],"reference_id":row["reference_id"],
                            "filename":row["filename"],"title":row["title"] or row["filename"],
                            "score":match["score"],"reasons":match["reasons"],
                            "fir_number":row.get("fir_number",""),"case_id":row.get("case_id","")})
    results.sort(key=lambda x:x["score"],reverse=True)
    return {"case_id":case_id,"matches":results,"count":len(results),
            "note":"Results are similarity-based investigative leads from uploaded synthetic case files."}

@app.post("/api/cctv/analyze")
async def analyze_cctv(
    file: UploadFile = File(...),
    case_id: str = Form("")
):
    if not file.filename or not file.filename.lower().endswith(".mp4"):
        raise HTTPException(status_code=415, detail="Upload an MP4 CCTV video.")
    if case_id.strip() and not get_case(case_id.strip()):
        raise HTTPException(status_code=404, detail="Case not found")
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Video file is empty.")
    if len(data) > 200 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Video is too large. Maximum size is 200 MB.")
    safe_name = re.sub(r"[^A-Za-z0-9._-]+","_",Path(file.filename).name)
    saved_name = uuid.uuid4().hex[:10] + "_" + safe_name
    saved_path = VIDEO_DIR / saved_name
    saved_path.write_bytes(data)

    if cv2 is None:
        return {"status":"UPLOADED","filename":file.filename,"video_url":f"/uploads/cctv/{saved_name}",
                "case_id":case_id,"analysis_available":False,
                "message":"Video uploaded, but OpenCV is not installed. Run: pip install opencv-python-headless",
                "observations":[]}

    cap=cv2.VideoCapture(str(saved_path))
    if not cap.isOpened():
        raise HTTPException(status_code=400, detail="The MP4 could not be opened as a video.")
    fps=float(cap.get(cv2.CAP_PROP_FPS) or 0)
    frames=int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
    width=int(cap.get(cv2.CAP_PROP_FRAME_WIDTH) or 0)
    height=int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT) or 0)
    duration=(frames/fps) if fps>0 else 0
    sample_every=max(1,int(fps*2)) if fps>0 else 30
    sample_limit=300
    sampled=0
    motion_events=0
    person_frames=0
    prev_gray=None
    hog=None
    try:
        hog=cv2.HOGDescriptor()
        hog.setSVMDetector(cv2.HOGDescriptor_getDefaultPeopleDetector())
    except Exception:
        hog=None
    observations=[]
    frame_index=0
    while frame_index < frames and sampled < sample_limit:
        cap.set(cv2.CAP_PROP_POS_FRAMES, frame_index)
        ok, frame=cap.read()
        if not ok: break
        gray=cv2.cvtColor(frame,cv2.COLOR_BGR2GRAY)
        motion_now=False
        if prev_gray is not None:
            diff=cv2.absdiff(gray,prev_gray)
            motion_ratio=float((diff>25).mean())
            motion_now = motion_ratio > 0.08
            if motion_now:
                motion_events += 1
        prev_gray=gray
        detected=0
        if hog is not None:
            try:
                boxes,_=hog.detectMultiScale(frame,winStride=(8,8),padding=(8,8),scale=1.05)
                detected=len(boxes)
            except Exception:
                detected=0
        if detected: person_frames += 1
        observations.append({
            "timestamp_seconds": round(frame_index/fps,2) if fps>0 else None,
            "frame": frame_index,
            "people_detected": detected,
            "movement_indicator": "MOTION" if motion_now else "OBSERVED"
        })
        sampled += 1
        frame_index += sample_every
    cap.release()
    # Keep response concise: return event samples with people/motion observations.
    events=[o for o in observations if o["people_detected"]>0 or o["movement_indicator"]=="MOTION"][:80]
    return {
        "status":"ANALYZED","filename":file.filename,"video_url":f"/uploads/cctv/{saved_name}",
        "case_id":case_id,"analysis_available":True,
        "video":{"fps":round(fps,2),"frames":frames,"duration_seconds":round(duration,2),
                 "width":width,"height":height,"sampled_frames":sampled},
        "summary":{"sampled_people_frames":person_frames,"movement_event_samples":motion_events,
                   "event_samples":len(events)},
        "observations":events,
        "method":"OpenCV frame sampling + anonymous person detection (HOG). No face identification is performed.",
        "note":"CCTV observations are visual leads for investigator review and do not identify a person or establish guilt."
    }

@app.post("/api/analyze-fir")
def analyze_fir(request: FIRAnalyzeRequest):
    if not request.text.strip():
        raise HTTPException(status_code=400, detail="FIR text cannot be empty")
    if request.priority.upper() not in {"LOW", "MEDIUM", "HIGH", "CRITICAL"}:
        raise HTTPException(status_code=400, detail="Invalid priority")

    graph = build_graph()
    result = analyze_fir_text(request.text, graph=graph)

    if request.create_case and result["case_id"] != "FIR-UNASSIGNED":
        existing = get_case(result["case_id"])
        if existing:
            result["case_created"] = False
            result["case_already_exists"] = True
        else:
            description = (
                f"Synthetic FIR {result.get('fir_number') or 'N/A'} analyzed. "
                f"Extracted {result['entity_count']} structured facts. "
                "See FIR Intelligence results for analytical leads."
            )
            create_case(
                result["case_id"],
                result.get("title") or "Synthetic FIR Investigation",
                description,
                request.priority.upper()
            )
            result["case_created"] = True

    return result


@app.post("/api/analyze-fir-file")
async def analyze_fir_file(
    file: UploadFile = File(...),
    create_case: bool = Form(False),
    priority: str = Form("MEDIUM")
):
    """Extract text from a synthetic FIR file and run the same FIR analysis pipeline."""
    if priority.upper() not in {"LOW", "MEDIUM", "HIGH", "CRITICAL"}:
        raise HTTPException(status_code=400, detail="Invalid priority")

    filename = (file.filename or "").lower()
    content_type = (file.content_type or "").lower()
    data = await file.read()

    if not data:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File is too large. Maximum size is 10 MB")

    try:
        if filename.endswith(".pdf") or content_type == "application/pdf":
            if PdfReader is None:
                raise HTTPException(status_code=500, detail="PDF support is not installed. Install pypdf.")
            from io import BytesIO
            reader = PdfReader(BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()

        elif filename.endswith(".docx") or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            if Document is None:
                raise HTTPException(status_code=500, detail="DOCX support is not installed. Install python-docx.")
            from io import BytesIO
            doc = Document(BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip()).strip()

        elif filename.endswith(".txt") or content_type.startswith("text/"):
            text = data.decode("utf-8-sig", errors="replace").strip()

        else:
            raise HTTPException(status_code=415, detail="Unsupported file type. Upload PDF, DOCX, or TXT.")

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read the document: {exc}")

    if not text:
        raise HTTPException(status_code=400, detail="No readable text was found in the document. Scanned/image-only PDFs need OCR support.")

    graph = build_graph()
    result = analyze_fir_text(text, graph=graph)
    result["uploaded_filename"] = file.filename
    result["document_type"] = "PDF" if filename.endswith(".pdf") else "DOCX" if filename.endswith(".docx") else "TXT"
    result["document_text_length"] = len(text)

    if create_case and result["case_id"] != "FIR-UNASSIGNED":
        existing = get_case(result["case_id"])
        if existing:
            result["case_created"] = False
            result["case_already_exists"] = True
        else:
            description = (
                f"Synthetic FIR {result.get('fir_number') or 'N/A'} uploaded as {file.filename}. "
                f"Extracted {result['entity_count']} structured facts. "
                "See FIR Intelligence results for analytical leads."
            )
            create_case(
                result["case_id"],
                result.get("title") or "Synthetic FIR Investigation",
                description,
                priority.upper()
            )
            result["case_created"] = True

    return result
